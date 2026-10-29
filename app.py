"""
Nexus | AI Second Brain
Streamlit + Supabase + Gemini.

Secrets (Streamlit Cloud > Settings > Secrets):
    GEMINI_API_KEY, SUPABASE_URL, SUPABASE_KEY      (required)
    USE_AUTH = "true"      (optional) per-user login with Supabase Auth
    APP_PASSWORD = "..."   (optional) simple password gate when USE_AUTH is off
    DAILY_GOAL = "3"       (optional) default daily review goal
"""
import streamlit as st
import pandas as pd
from datetime import datetime, date, timedelta
import os
import re
import json
import hmac
from io import BytesIO
import PyPDF2
from PIL import Image
from pptx import Presentation
from pptx.util import Pt
from google import genai
from google.genai import types

# --- Optional imports (the app degrades gracefully without them) ---
try:
    from supabase import create_client
    SUPABASE_INSTALLED = True
except ImportError:
    SUPABASE_INSTALLED = False

try:
    from streamlit_agraph import agraph, Node, Edge, Config
    AGRAPH_INSTALLED = True
except ImportError:
    AGRAPH_INSTALLED = False

try:
    import requests
    from bs4 import BeautifulSoup
    WEB_OK = True
except ImportError:
    WEB_OK = False

try:
    from youtube_transcript_api import YouTubeTranscriptApi
    YT_OK = True
except ImportError:
    YT_OK = False

st.set_page_config(page_title="Nexus | AI Second Brain", page_icon="🧠", layout="wide")

st.markdown("""
<style>
    @import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700;800&display=swap');
    :root {
        --nx-bg: #0b1020; --nx-border: rgba(148, 163, 184, .18);
        --nx-text: #e8eefb; --nx-muted: #9ba9c4;
    }
    html, body, [class*="css"] { font-family: 'Inter', sans-serif; }
    .stApp {
        background:
            radial-gradient(ellipse at 12% 0%, rgba(69, 96, 210, .18), transparent 34%),
            radial-gradient(ellipse at 92% 8%, rgba(27, 196, 191, .10), transparent 28%),
            #0b1020;
        color: var(--nx-text);
    }
    [data-testid="stHeader"] { background: rgba(11, 16, 32, .75); }
    [data-testid="stSidebar"] {
        background: linear-gradient(180deg, #10182c 0%, #0b1020 100%);
        border-right: 1px solid var(--nx-border);
    }
    [data-testid="stSidebar"] > div { padding-top: 1.25rem; }
    .block-container { padding-top: 1.5rem; padding-bottom: 3rem; max-width: 1600px; }
    .main-header {
        position: relative; overflow: hidden;
        background: linear-gradient(118deg, #18284e 0%, #263c78 48%, #087e91 100%);
        padding: 2rem 2.1rem; border: 1px solid rgba(255,255,255,.13);
        border-radius: 24px; color: white; margin-bottom: 1.35rem;
        box-shadow: 0 18px 55px rgba(0,0,0,.22);
    }
    .main-header:after {
        content: ""; position: absolute; width: 270px; height: 270px;
        right: -75px; top: -145px; border: 1px solid rgba(255,255,255,.16);
        border-radius: 50%;
        box-shadow: 0 0 0 30px rgba(255,255,255,.035), 0 0 0 65px rgba(255,255,255,.025);
        pointer-events: none;
    }
    .eyebrow { text-transform: uppercase; letter-spacing: .18em; font-size: .72rem; color: #b9f5ee; font-weight: 800; margin-bottom: .55rem; }
    .main-header h1 { color: #fff; margin: 0; font-size: clamp(1.7rem, 3vw, 2.55rem); font-weight: 800; letter-spacing: -.045em; }
    .main-header p { color: #dce8ff; margin: .65rem 0 0; font-size: .98rem; max-width: 760px; line-height: 1.65; }
    .status-badge {
        padding: 9px 11px; border-radius: 12px; font-size: .76rem; font-weight: 700;
        text-align: center; margin-bottom: 8px; border: 1px solid var(--nx-border);
        background: rgba(255,255,255,.04);
    }
    .badge-ok { color: #5eead4; border-color: rgba(45,212,191,.24); background: rgba(45,212,191,.08); }
    .badge-warn { color: #fbbf24; border-color: rgba(251,191,36,.24); background: rgba(251,191,36,.08); }
    div[data-testid="stMetric"] {
        background: linear-gradient(145deg, rgba(23,34,59,.95), rgba(16,24,43,.78));
        border: 1px solid var(--nx-border); padding: 1rem 1.1rem;
        border-radius: 18px; box-shadow: 0 8px 28px rgba(0,0,0,.10);
    }
    div[data-testid="stMetricLabel"] { color: #aab8d4; font-weight: 600; }
    div[data-testid="stMetricValue"] { color: #f4f7ff; font-weight: 800; }
    [data-testid="stVerticalBlockBorderWrapper"] > div {
        border-color: var(--nx-border) !important; border-radius: 18px !important;
        background: rgba(18,27,48,.42);
    }
    .stTabs [data-baseweb="tab-list"] {
        gap: 7px; border-bottom: 1px solid var(--nx-border);
        background: rgba(255,255,255,.025); padding: 7px; border-radius: 16px;
    }
    .stTabs [data-baseweb="tab"] {
        height: 43px; border-radius: 11px; padding: 0 15px;
        color: #aab8d4; font-weight: 650; border: 1px solid transparent;
    }
    .stTabs [aria-selected="true"] {
        background: linear-gradient(120deg, rgba(124,156,255,.22), rgba(77,226,208,.10));
        color: #f5f8ff !important; border-color: rgba(124,156,255,.22);
    }
    .stButton > button, .stDownloadButton > button {
        width: 100%; border-radius: 12px; min-height: 2.65rem; font-weight: 700;
        transition: transform .15s ease, box-shadow .15s ease, border-color .15s ease;
        border: 1px solid var(--nx-border);
    }
    .stButton > button:hover, .stDownloadButton > button:hover {
        transform: translateY(-1px); box-shadow: 0 8px 24px rgba(0,0,0,.18);
        border-color: rgba(124,156,255,.55);
    }
    .stButton > button[kind="primary"] { background: linear-gradient(115deg, #647ff5, #0eaaa8); color: white; border: 0; }
    input, textarea, [data-baseweb="select"] > div { border-radius: 11px !important; }
    .flashcard {
        background: linear-gradient(135deg, #1a2c51 0%, #111b31 100%);
        color: #f8fafc; padding: 1.2rem; border-radius: 16px;
        border: 1px solid rgba(124,156,255,.22); border-left: 4px solid #4de2d0;
        margin-bottom: .8rem; box-shadow: 0 8px 25px rgba(0,0,0,.12);
    }
    .empty-state {
        padding: 2rem; text-align: center; border-radius: 18px;
        border: 1px dashed rgba(148,163,184,.3); background: rgba(255,255,255,.025); color: #aab8d4;
    }
    hr { border-color: var(--nx-border) !important; }
    @media (max-width: 800px) {
        .block-container { padding: 1rem .8rem 2rem; }
        .main-header { padding: 1.4rem; border-radius: 18px; }
        .stTabs [data-baseweb="tab"] { padding: 0 9px; font-size: .78rem; }
    }
</style>
""", unsafe_allow_html=True)

# ================= CONSTANTS =================
CAT_OPTIONS = ["Project Management (PM)", "Technical / Engineering", "AI & Automation", "Business & Strategy"]
STAT_OPTIONS = ["Not Started", "In Progress", "Completed"]

TOPIC_DEFAULTS = {
    "title": "", "category": "General", "status": "In Progress", "mastery": 50,
    "notes": "", "ai_summary": "", "related_topics": "", "prereq_topics": "", "tags": "",
    "ease": 2.5, "interval_days": 0, "repetitions": 0, "next_review": "", "updated_at": "",
}
TOPIC_COLS = list(TOPIC_DEFAULTS.keys())


# ================= SECRETS & CLIENTS =================
def get_secret(key):
    if key in os.environ and os.environ[key]:
        return os.environ[key].strip()
    try:
        if key in st.secrets:
            return str(st.secrets[key]).strip()
    except Exception:
        pass
    return None


def sanitize_supabase_url(raw_url):
    """Fix dashboard URLs / extra paths into a valid https://<ref>.supabase.co URL"""
    if not raw_url:
        return None
    raw_url = raw_url.strip().strip('"').strip("'")
    m = re.search(r'supabase\.com/dashboard/project/([a-z0-9]+)', raw_url)
    if m:
        return f"https://{m.group(1)}.supabase.co"
    m = re.search(r'(https://[a-z0-9]+\.supabase\.co)', raw_url)
    if m:
        return m.group(1)
    if re.match(r'^[a-z0-9]{15,25}$', raw_url):
        return f"https://{raw_url}.supabase.co"
    return raw_url.rstrip('/')


GEMINI_API_KEY = get_secret("GEMINI_API_KEY")
SUPABASE_URL = sanitize_supabase_url(get_secret("SUPABASE_URL"))
SUPABASE_KEY = get_secret("SUPABASE_KEY")
APP_PASSWORD = get_secret("APP_PASSWORD")
USE_AUTH = (get_secret("USE_AUTH") or "").lower() in ("1", "true", "yes")
try:
    DEFAULT_GOAL = int(get_secret("DAILY_GOAL") or 3)
except ValueError:
    DEFAULT_GOAL = 3

# Conflict targets for upserts (they must match the unique indexes in the SQL setup)
if USE_AUTH:
    CONFLICT = {"topics": "user_id,title", "glossary": "user_id,term", "flashcards": "user_id,topic,q"}
else:
    CONFLICT = {"topics": "title", "glossary": "term", "flashcards": "topic,q"}

# --- Simple password gate (only when per-user auth is off) ---
if APP_PASSWORD and not USE_AUTH and not st.session_state.get("authed"):
    st.markdown("### 🔒 Nexus")
    pw = st.text_input("Password", type="password")
    if pw:
        if hmac.compare_digest(pw, APP_PASSWORD):
            st.session_state.authed = True
            st.rerun()
        else:
            st.error("Wrong password")
    st.stop()


@st.cache_resource(show_spinner=False)
def get_ai_client(key):
    return genai.Client(api_key=key) if key else None


@st.cache_resource(show_spinner=False)
def get_shared_client(url, key):
    if SUPABASE_INSTALLED and url and key:
        return create_client(url, key)
    return None


ai_client = get_ai_client(GEMINI_API_KEY)
shared_client, _client_err = None, ""
if not USE_AUTH:
    try:
        shared_client = get_shared_client(SUPABASE_URL, SUPABASE_KEY)
    except Exception as e:
        _client_err = str(e)


def db():
    """Supabase client. With auth on, each browser session has its own signed-in client."""
    if USE_AUTH:
        return st.session_state.get("sb_client")
    return shared_client


# --- Per-user sign in (Supabase Auth) ---
if USE_AUTH and not st.session_state.get("sb_user"):
    st.markdown("### 🔐 Nexus sign in")
    if not (SUPABASE_INSTALLED and SUPABASE_URL and SUPABASE_KEY):
        st.error("SUPABASE_URL / SUPABASE_KEY secrets are required when USE_AUTH is on.")
        st.stop()
    auth_mode = st.radio("Mode", ["Sign in", "Sign up"], horizontal=True, label_visibility="collapsed")
    a_email = st.text_input("Email")
    a_pw = st.text_input("Password", type="password")
    logged_in = False
    if st.button("Continue", type="primary"):
        try:
            c = create_client(SUPABASE_URL, SUPABASE_KEY)
            if auth_mode == "Sign up":
                c.auth.sign_up({"email": a_email, "password": a_pw})
                st.success("Account created. If email confirmation is enabled in Supabase, confirm your email, then sign in.")
            else:
                res = c.auth.sign_in_with_password({"email": a_email, "password": a_pw})
                st.session_state.sb_client = c
                st.session_state.sb_user = (res.user.email if res and res.user else a_email)
                logged_in = True
        except Exception as e:
            st.error(f"Auth failed: {str(e)[:200]}")
    if logged_in:
        st.rerun()
    st.stop()


# ================= SMALL HELPERS =================
def today_str():
    return date.today().isoformat()


def split_csv(value):
    return [x.strip() for x in str(value or "").split(",") if x.strip()]


def uniq(items):
    seen, out = set(), []
    for x in items:
        k = str(x).lower()
        if k not in seen:
            seen.add(k)
            out.append(x)
    return out


def extract_links(text):
    return uniq([m.strip() for m in re.findall(r"\[\[([^\]]+)\]\]", text or "")])


def render_links(text):
    """Turn [[Topic]] into bold links text: existing topics get a link icon, missing ones are marked."""
    known = {t["title"].lower() for t in st.session_state.topics}

    def _sub(m):
        name = m.group(1).strip()
        return f"**🔗 {name}**" if name.lower() in known else f"*{name} (no such topic)*"
    return re.sub(r"\[\[([^\]]+)\]\]", _sub, text or "")


def backlinks_of(title):
    """Topics whose notes mention [[title]]."""
    out = []
    for t in st.session_state.topics:
        if t["title"].lower() == title.lower():
            continue
        if title.lower() in [x.lower() for x in extract_links(t.get("notes"))]:
            out.append(t["title"])
    return out


def snippet(text, query, width=90):
    text = text or ""
    i = text.lower().find(query.lower())
    if i < 0:
        return text[:width * 2]
    a, b = max(0, i - width), min(len(text), i + len(query) + width)
    return ("..." if a > 0 else "") + text[a:b].replace("\n", " ") + ("..." if b < len(text) else "")


# ================= DATABASE LAYER =================
def explain_db_error(e):
    """Turn raw Supabase/Postgres errors into an actionable message."""
    s = str(e)
    low = s.lower()
    if "row-level security" in low or "42501" in low:
        hint = "Row Level Security is on but no policy allows this. Run the setup SQL from the Cloud Diagnostics tab."
    elif "23502" in low and "user_id" in low:
        hint = "user_id is empty. Run the AUTH SQL from the Cloud Diagnostics tab and make sure you are signed in."
    elif "42p10" in low or "no unique or exclusion constraint" in low:
        hint = "A UNIQUE index is missing, so upsert fails. Run the setup SQL from the Cloud Diagnostics tab."
    elif "pgrst205" in low or "could not find the table" in low or "42p01" in low:
        hint = "A table does not exist in Supabase. Run the setup SQL in the SQL Editor."
    elif "pgrst204" in low or ("column" in low and "does not exist" in low):
        hint = "A column is missing. Run the setup SQL (section 2) from the Cloud Diagnostics tab."
    elif "invalid api key" in low or "jwt" in low or "apikey" in low:
        hint = "The key or session is invalid. Check SUPABASE_KEY, or sign in again."
    elif "name or service not known" in low or "connection" in low or "timed out" in low:
        hint = "Cannot reach Supabase. Check SUPABASE_URL (or whether the project is paused)."
    else:
        hint = "Unexpected error."
    return f"{hint}\n\nDetails: {s[:300]}"


def run_db_check():
    """Real read + write probe on all four tables. Returns (ok, message)."""
    if not SUPABASE_INSTALLED:
        return False, "The supabase package is not installed (check requirements.txt)."
    if _client_err:
        return False, explain_db_error(_client_err)
    c = db()
    if not c:
        return False, "SUPABASE_URL / SUPABASE_KEY secrets are missing: data is local only and is lost on every refresh."
    try:
        for tbl, col in (("topics", "title"), ("glossary", "term"), ("flashcards", "q"), ("activity_log", "topic")):
            c.table(tbl).select(col).limit(1).execute()
        probe = dict(TOPIC_DEFAULTS, title="__nexus_probe__", updated_at=today_str())
        c.table("topics").upsert(probe, on_conflict=CONFLICT["topics"]).execute()
        c.table("topics").delete().eq("title", "__nexus_probe__").execute()
        c.table("glossary").upsert({"term": "__nexus_probe__", "category": "x", "context": "x"},
                                   on_conflict=CONFLICT["glossary"]).execute()
        c.table("glossary").delete().eq("term", "__nexus_probe__").execute()
        c.table("flashcards").upsert({"topic": "__nexus_probe__", "q": "x", "a": "x", "created_at": today_str()},
                                     on_conflict=CONFLICT["flashcards"]).execute()
        c.table("flashcards").delete().eq("topic", "__nexus_probe__").execute()
        c.table("activity_log").insert({"topic": "__nexus_probe__", "kind": "probe", "value": 0,
                                        "logged_on": today_str()}).execute()
        c.table("activity_log").delete().eq("topic", "__nexus_probe__").execute()
        return True, ""
    except Exception as e:
        return False, explain_db_error(e)


def fetch_all(table):
    """Read a whole table (Supabase returns max 1000 rows per request, so page through)."""
    rows, start = [], 0
    while True:
        batch = db().table(table).select("*").order("id").range(start, start + 999).execute().data or []
        rows += batch
        if len(batch) < 1000:
            break
        start += 1000
    return rows


def norm_topic(r):
    t = {}
    for k, default in TOPIC_DEFAULTS.items():
        v = r.get(k)
        t[k] = default if v is None else v
    t["mastery"] = int(t["mastery"] or 0)
    t["ease"] = float(t["ease"] or 2.5)
    t["interval_days"] = int(t["interval_days"] or 0)
    t["repetitions"] = int(t["repetitions"] or 0)
    return t


def load_data_from_cloud():
    if not db():
        return False
    try:
        topics = [norm_topic(r) for r in fetch_all("topics")]
        topics.sort(key=lambda t: t.get("updated_at") or "", reverse=True)
        st.session_state.topics = topics
        st.session_state.glossary = fetch_all("glossary")
        st.session_state.flashcards = fetch_all("flashcards")
        st.session_state.activity = fetch_all("activity_log")
        return True
    except Exception as e:
        st.session_state.db_error_msg = explain_db_error(e)
        return False


if "data_initialized" not in st.session_state:
    st.session_state.topics = []
    st.session_state.glossary = []
    st.session_state.flashcards = []
    st.session_state.activity = []
    ok_, msg_ = run_db_check()
    st.session_state.db_verified = ok_
    st.session_state.db_error_msg = msg_
    if ok_:
        load_data_from_cloud()
    st.session_state.last_save_error = ""
    st.session_state.data_initialized = True


def _cloud_write(table, item, conflict_col):
    """Upsert (dict or list of dicts). Returns '' on success or an error message."""
    c = db()
    if not c:
        return "Local mode: data is not saved to the cloud."
    try:
        c.table(table).upsert(item, on_conflict=conflict_col).execute()
        st.session_state.db_verified = True
        return ""
    except Exception as e:
        st.session_state.db_verified = False
        return explain_db_error(e)


def _cloud_insert(table, item):
    c = db()
    if not c:
        return "Local mode: data is not saved to the cloud."
    try:
        c.table(table).insert(item).execute()
        return ""
    except Exception as e:
        return explain_db_error(e)


def _cloud_delete(table, **match):
    c = db()
    if not c:
        return ""
    try:
        q = c.table(table).delete()
        for k, v in match.items():
            q = q.eq(k, v)
        q.execute()
        return ""
    except Exception as e:
        return explain_db_error(e)


def log_activity(topic, kind, value):
    """kind: mastery | review | quiz | flashcard"""
    item = {"topic": topic, "kind": kind, "value": int(value), "logged_on": today_str()}
    st.session_state.activity.append(item)
    return _cloud_insert("activity_log", item)


def upsert_topic(title, **fields):
    """Create or update a topic. Only the fields you pass are changed."""
    clean = title.strip()
    idx = next((i for i, t in enumerate(st.session_state.topics) if t["title"].lower() == clean.lower()), None)
    curr = st.session_state.topics[idx] if idx is not None else dict(TOPIC_DEFAULTS)
    item = {c: curr.get(c, TOPIC_DEFAULTS[c]) for c in TOPIC_COLS}
    item["title"] = curr["title"] if idx is not None else clean
    for k, v in fields.items():
        if v is not None and k in TOPIC_DEFAULTS and k != "title":
            item[k] = v
    item["mastery"] = max(0, min(100, int(item["mastery"])))
    item["updated_at"] = today_str()
    old_m = curr.get("mastery") if idx is not None else None

    if idx is not None:
        st.session_state.topics[idx] = item
        action = "updated"
    else:
        st.session_state.topics.insert(0, item)
        action = "created"

    err = _cloud_write("topics", item, CONFLICT["topics"])
    if old_m is None or old_m != item["mastery"]:
        err = err or log_activity(item["title"], "mastery", item["mastery"])
    st.session_state.last_save_error = err
    return action, err


def delete_topic(title):
    st.session_state.topics = [t for t in st.session_state.topics if t["title"] != title]
    st.session_state.last_save_error = _cloud_delete("topics", title=title)


def upsert_glossary(term, category, context):
    clean = term.strip()
    item = {"term": clean, "category": category, "context": context}
    idx = next((i for i, g in enumerate(st.session_state.glossary) if g["term"].lower() == clean.lower()), None)
    if idx is not None:
        st.session_state.glossary[idx] = item
    else:
        st.session_state.glossary.insert(0, item)
    err = _cloud_write("glossary", item, CONFLICT["glossary"])
    if err:
        st.session_state.last_save_error = err
    return err


def delete_glossary(term):
    st.session_state.glossary = [g for g in st.session_state.glossary if g["term"] != term]
    err = _cloud_delete("glossary", term=term)
    if err:
        st.session_state.last_save_error = err


def add_flashcards(topic, cards):
    """Save cards (list of {'q','a'}). Duplicates (same topic + question) are updated."""
    rows = []
    for c in cards:
        q, a = str(c.get("q", "")).strip(), str(c.get("a", "")).strip()
        if not q:
            continue
        item = {"topic": topic, "q": q, "a": a, "created_at": today_str()}
        existing = next((i for i, x in enumerate(st.session_state.flashcards)
                         if x["topic"] == topic and x["q"] == q), None)
        if existing is not None:
            st.session_state.flashcards[existing] = item
        else:
            st.session_state.flashcards.append(item)
        rows.append(item)
    if not rows:
        return ""
    err = _cloud_write("flashcards", rows, CONFLICT["flashcards"])
    if err:
        st.session_state.last_save_error = err
    return err


def delete_flashcard(topic, q):
    st.session_state.flashcards = [x for x in st.session_state.flashcards
                                   if not (x["topic"] == topic and x["q"] == q)]
    err = _cloud_delete("flashcards", topic=topic, q=q)
    if err:
        st.session_state.last_save_error = err


# ================= STUDY LOGIC =================
GRADES = {"Again": 1, "Hard": 3, "Good": 4, "Easy": 5}
MASTERY_DELTA = {1: -10, 3: 0, 4: 5, 5: 10}


def sm2_next(t, grade):
    """SM-2 spaced repetition. grade: 1 (again) .. 5 (easy)."""
    ease = float(t.get("ease") or 2.5)
    reps = int(t.get("repetitions") or 0)
    interval = int(t.get("interval_days") or 0)
    if grade < 3:
        reps, interval = 0, 1
    else:
        reps += 1
        interval = 1 if reps == 1 else (6 if reps == 2 else max(1, round(interval * ease)))
    ease = max(1.3, ease + 0.1 - (5 - grade) * (0.08 + (5 - grade) * 0.02))
    return {"ease": round(ease, 2), "interval_days": interval, "repetitions": reps,
            "next_review": (date.today() + timedelta(days=interval)).isoformat()}


def due_topics():
    today = today_str()
    out = [t for t in st.session_state.topics
           if t.get("status") != "Not Started" and (not t.get("next_review") or t["next_review"] <= today)]
    out.sort(key=lambda t: (int(t.get("mastery") or 0), t.get("next_review") or ""))
    return out


def study_days():
    return {a["logged_on"] for a in st.session_state.activity if a.get("kind") in ("review", "quiz", "flashcard")}


def compute_streak():
    days = study_days()
    d = date.today()
    if d.isoformat() not in days:
        d -= timedelta(days=1)
    n = 0
    while d.isoformat() in days:
        n += 1
        d -= timedelta(days=1)
    return n


def reviews_today():
    t = today_str()
    return sum(1 for a in st.session_state.activity if a.get("kind") == "review" and a.get("logged_on") == t)


def learning_path(topics):
    """Order topics so prerequisites come first; among ready topics the weakest goes first."""
    by = {t["title"]: t for t in topics}
    pre = {k: [p for p in split_csv(v.get("prereq_topics")) if p in by and p != k] for k, v in by.items()}
    done, order, remaining, cyclic = set(), [], set(by), set()
    m = lambda k: int(by[k].get("mastery") or 0)
    while remaining:
        ready = [k for k in remaining if all(p in done for p in pre[k])]
        if not ready:
            pick = min(remaining, key=lambda k: (m(k), k))
            cyclic.add(pick)
        else:
            pick = min(ready, key=lambda k: (m(k), k))
        remaining.discard(pick)
        done.add(pick)
        order.append(pick)
    rows = []
    for i, k in enumerate(order, 1):
        note = "circular prerequisite" if k in cyclic else ("mastered" if m(k) >= 80 else "needs work" if m(k) < 40 else "in progress")
        rows.append({"#": i, "topic": k, "mastery": m(k), "status": by[k].get("status"),
                     "after": ", ".join(pre[k]), "note": note})
    return rows


# ================= AI HELPERS =================
def parse_json_loose(text):
    """Extract JSON from a model reply even if wrapped in ``` fences or extra prose."""
    t = text.strip()
    t = re.sub(r"^```(?:json)?", "", t).strip()
    t = re.sub(r"```$", "", t).strip()
    try:
        return json.loads(t)
    except Exception:
        pass
    starts = [p for p in (t.find("["), t.find("{")) if p != -1]
    if not starts:
        raise ValueError("No JSON found")
    start = min(starts)
    end = max(t.rfind("]"), t.rfind("}"))
    return json.loads(t[start:end + 1])


def normalize_links(data):
    """Return {topic: {'related': [...], 'prereqs': [...]}} from several possible JSON shapes."""
    def as_list(v):
        if isinstance(v, list):
            return [str(x).strip() for x in v if str(x).strip()]
        return split_csv(v)
    out = {}
    if isinstance(data, dict):
        for k, v in data.items():
            if isinstance(v, dict):
                out[str(k)] = {"related": as_list(v.get("related") or v.get("related_topics")),
                               "prereqs": as_list(v.get("prereqs") or v.get("prerequisites"))}
            else:
                out[str(k)] = {"related": as_list(v), "prereqs": []}
    elif isinstance(data, list):
        for row in data:
            if isinstance(row, dict):
                k = row.get("topic") or row.get("title") or row.get("source")
                if k:
                    out[str(k)] = {"related": as_list(row.get("related") or row.get("related_topics") or row.get("targets")),
                                   "prereqs": as_list(row.get("prereqs") or row.get("prerequisites"))}
    return out


def ask_gemini(prompt, media=None, selected_model="gemini-3.8-flash", json_mode=False):
    """media: list of PIL images and/or google.genai types.Part objects (pdf / audio bytes)."""
    if not ai_client:
        return False, "⚠️ AI Engine offline. Please verify GEMINI_API_KEY."
    contents = [prompt] + list(media or [])
    config = types.GenerateContentConfig(response_mime_type="application/json") if json_mode else None
    chain = []
    for m in [selected_model, "gemini-3.8-flash", "gemini-3.6-flash", "gemini-3.5-flash-lite"]:
        if m and m not in chain:
            chain.append(m)
    errors = []
    for m in chain:
        try:
            response = ai_client.models.generate_content(model=m, contents=contents, config=config)
            if response and response.text:
                return True, response.text
            errors.append(f"{m}: empty response")
        except Exception as e:
            errors.append(f"{m}: {str(e)[:160]}")
    return False, "⚠️ AI request failed on every model tried:\n\n" + "\n\n".join(errors)


def fetch_url_text(url):
    """Return (ok, text). Handles normal web pages and YouTube links (transcript)."""
    url = (url or "").strip()
    if not url:
        return False, "Empty URL."
    yt = re.search(r"(?:v=|youtu\.be/|shorts/|embed/)([A-Za-z0-9_-]{11})", url)
    if yt:
        if not YT_OK:
            return False, "youtube-transcript-api is not installed (add it to requirements.txt)."
        vid = yt.group(1)
        try:
            try:
                fetched = YouTubeTranscriptApi().fetch(vid, languages=["en", "ar"])
                text = " ".join(sn.text for sn in fetched)
            except AttributeError:
                data = YouTubeTranscriptApi.get_transcript(vid, languages=["en", "ar"])
                text = " ".join(d["text"] for d in data)
            return True, text[:12000]
        except Exception as e:
            return False, f"Could not get the transcript (YouTube may block cloud servers): {str(e)[:150]}"
    if not WEB_OK:
        return False, "requests / beautifulsoup4 are not installed (add them to requirements.txt)."
    try:
        r = requests.get(url, timeout=15, headers={"User-Agent": "Mozilla/5.0 (Nexus second brain)"})
        r.raise_for_status()
        soup = BeautifulSoup(r.text, "html.parser")
        for tag in soup(["script", "style", "nav", "footer", "header", "aside", "form"]):
            tag.decompose()
        text = re.sub(r"\s+", " ", soup.get_text(" ", strip=True))
        return (True, text[:12000]) if text else (False, "No readable text found on that page.")
    except Exception as e:
        return False, f"Could not open the page: {str(e)[:150]}"


def pdf_to_input(uploaded):
    """Return (text, media_part). Uses text extraction; falls back to sending the PDF to Gemini (OCR)."""
    data = uploaded.getvalue()
    text = ""
    try:
        reader = PyPDF2.PdfReader(BytesIO(data))
        for page in reader.pages[:10]:
            text += (page.extract_text() or "") + "\n"
    except Exception:
        pass
    if len(text.strip()) > 80:
        return text, None
    return "", types.Part.from_bytes(data=data, mime_type="application/pdf")


def build_dot(topics, show_hubs=True, highlight_weak=False):
    """Static Graphviz graph (no extra library needed)."""
    def esc(x):
        return str(x).replace('"', "'")
    titles = {t["title"] for t in topics}
    lines = ["graph G {", 'bgcolor="transparent";',
             'node [style=filled, fontcolor="white", fontname="Helvetica", color="#334155"];',
             'edge [color="#64748b"];']
    if show_hubs:
        for c in set(t.get("category") or "General" for t in topics):
            lines.append(f'"HUB_{esc(c)}" [label="{esc(c)}", shape=hexagon, fillcolor="#1e3a8a"];')
    seen = set()
    for t in topics:
        mast = int(t.get("mastery") or 50)
        lines.append(f'"{esc(t["title"])}" [label="{esc(t["title"])}\\n({mast}%)", fillcolor="{node_color(t, highlight_weak)}"];')
        if show_hubs:
            lines.append(f'"HUB_{esc(t.get("category") or "General")}" -- "{esc(t["title"])}";')
    for t in topics:
        for r in split_csv(t.get("related_topics")):
            tgt = next((x for x in titles if x.lower() == r.lower()), None)
            key = ("rel", frozenset((t["title"], tgt or "")))
            if tgt and tgt != t["title"] and key not in seen:
                seen.add(key)
                lines.append(f'"{esc(t["title"])}" -- "{esc(tgt)}" [color="#38bdf8", penwidth=2];')
        for p in split_csv(t.get("prereq_topics")):
            src = next((x for x in titles if x.lower() == p.lower()), None)
            if src and src != t["title"]:
                lines.append(f'"{esc(src)}" -- "{esc(t["title"])}" [dir=forward, color="#f59e0b", penwidth=2];')
    lines.append("}")
    return "\n".join(lines)


def node_color(t, highlight_weak=False):
    mast = int(t.get("mastery") or 50)
    if highlight_weak and mast >= 40:
        return "#475569"
    if t.get("status") == "Completed" or mast >= 80:
        return "#10b981"
    return "#f59e0b" if mast >= 40 else "#ef4444"



# ================= SQL TEXT SHOWN IN THE DIAGNOSTICS TAB =================
SQL_SETUP = """
-- 1) TABLES
create table if not exists topics (
  id bigint generated always as identity primary key,
  title text not null,
  category text, status text, mastery int default 50,
  notes text default '', ai_summary text default '',
  related_topics text default '', prereq_topics text default '', tags text default '',
  ease real default 2.5, interval_days int default 0, repetitions int default 0,
  next_review text default '', updated_at text
);
create table if not exists glossary (
  id bigint generated always as identity primary key,
  term text not null, category text, context text
);
create table if not exists flashcards (
  id bigint generated always as identity primary key,
  topic text not null, q text not null, a text, created_at text
);
create table if not exists activity_log (
  id bigint generated always as identity primary key,
  topic text, kind text, value int, logged_on text
);

-- 2) COLUMNS (for tables created by an older version)
alter table topics add column if not exists category text;
alter table topics add column if not exists status text;
alter table topics add column if not exists mastery int default 50;
alter table topics add column if not exists notes text default '';
alter table topics add column if not exists ai_summary text default '';
alter table topics add column if not exists related_topics text default '';
alter table topics add column if not exists prereq_topics text default '';
alter table topics add column if not exists tags text default '';
alter table topics add column if not exists ease real default 2.5;
alter table topics add column if not exists interval_days int default 0;
alter table topics add column if not exists repetitions int default 0;
alter table topics add column if not exists next_review text default '';
alter table topics add column if not exists updated_at text;
alter table glossary add column if not exists category text;
alter table glossary add column if not exists context text;

-- 3) REMOVE DUPLICATES + UNIQUE INDEXES (needed by upsert)
delete from topics a using topics b where a.ctid < b.ctid and a.title = b.title;
delete from glossary a using glossary b where a.ctid < b.ctid and a.term = b.term;
delete from flashcards a using flashcards b where a.ctid < b.ctid and a.topic = b.topic and a.q = b.q;
create unique index if not exists topics_title_uidx on topics (title);
create unique index if not exists glossary_term_uidx on glossary (term);
create unique index if not exists flashcards_topic_q_uidx on flashcards (topic, q);

-- 4) ROW LEVEL SECURITY + OPEN POLICIES (shared mode)
alter table topics enable row level security;
alter table glossary enable row level security;
alter table flashcards enable row level security;
alter table activity_log enable row level security;
drop policy if exists "nexus_all_topics" on topics;
drop policy if exists "nexus_all_glossary" on glossary;
drop policy if exists "nexus_all_flashcards" on flashcards;
drop policy if exists "nexus_all_activity" on activity_log;
create policy "nexus_all_topics" on topics for all to anon, authenticated using (true) with check (true);
create policy "nexus_all_glossary" on glossary for all to anon, authenticated using (true) with check (true);
create policy "nexus_all_flashcards" on flashcards for all to anon, authenticated using (true) with check (true);
create policy "nexus_all_activity" on activity_log for all to anon, authenticated using (true) with check (true);
"""

SQL_AUTH = """
-- Run the main setup first. This switches every table to per-user data.
alter table topics add column if not exists user_id uuid default auth.uid();
alter table glossary add column if not exists user_id uuid default auth.uid();
alter table flashcards add column if not exists user_id uuid default auth.uid();
alter table activity_log add column if not exists user_id uuid default auth.uid();

drop index if exists topics_title_uidx;
drop index if exists glossary_term_uidx;
drop index if exists flashcards_topic_q_uidx;
create unique index if not exists topics_user_title_uidx on topics (user_id, title);
create unique index if not exists glossary_user_term_uidx on glossary (user_id, term);
create unique index if not exists flashcards_user_topic_q_uidx on flashcards (user_id, topic, q);

drop policy if exists "nexus_all_topics" on topics;
drop policy if exists "nexus_all_glossary" on glossary;
drop policy if exists "nexus_all_flashcards" on flashcards;
drop policy if exists "nexus_all_activity" on activity_log;
drop policy if exists "nexus_own_topics" on topics;
drop policy if exists "nexus_own_glossary" on glossary;
drop policy if exists "nexus_own_flashcards" on flashcards;
drop policy if exists "nexus_own_activity" on activity_log;
create policy "nexus_own_topics" on topics for all to authenticated using (user_id = auth.uid()) with check (user_id = auth.uid());
create policy "nexus_own_glossary" on glossary for all to authenticated using (user_id = auth.uid()) with check (user_id = auth.uid());
create policy "nexus_own_flashcards" on flashcards for all to authenticated using (user_id = auth.uid()) with check (user_id = auth.uid());
create policy "nexus_own_activity" on activity_log for all to authenticated using (user_id = auth.uid()) with check (user_id = auth.uid());

-- Claim your existing rows: sign up in the app first, copy your user id from
-- Supabase > Authentication > Users, then run (replace the id) for each table:
-- update topics set user_id = 'YOUR-USER-ID' where user_id is null;
-- update glossary set user_id = 'YOUR-USER-ID' where user_id is null;
-- update flashcards set user_id = 'YOUR-USER-ID' where user_id is null;
-- update activity_log set user_id = 'YOUR-USER-ID' where user_id is null;
"""

def build_advanced_pptx(topic, audience, slides_json):
    prs = Presentation()
    s0 = prs.slides.add_slide(prs.slide_layouts[0])
    s0.shapes.title.text = topic
    s0.placeholders[1].text = f"Target Audience: {audience}\nGenerated by Nexus AI Engine"
    for s in slides_json:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = s.get("title", "Slide")
        tf = slide.placeholders[1].text_frame
        tf.word_wrap = True
        for idx, b in enumerate(s.get("bullets", [])):
            p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
            p.text = str(b)
            p.font.size = Pt(20)
        slide.notes_slide.notes_text_frame.text = (
            f"⏱️ Time: {s.get('time', '2 mins')}\n\n🎙️ Script:\n{s.get('speaker_notes', '')}")
    stream = BytesIO()
    prs.save(stream)
    stream.seek(0)
    return stream


# ================= SIDEBAR =================
topic_names = [t["title"] for t in st.session_state.topics]

with st.sidebar:
    st.markdown("### ⚡ Nexus Control Center")
    if USE_AUTH:
        st.caption(f"Signed in as {st.session_state.get('sb_user')}")
        if st.button("Sign out"):
            try:
                db().auth.sign_out()
            except Exception:
                pass
            for k in list(st.session_state.keys()):
                del st.session_state[k]
            st.rerun()

    ai_badge = ('<div class="status-badge badge-ok">🤖 AI: Ready</div>' if ai_client
                else '<div class="status-badge badge-warn">🤖 AI: Missing Key</div>')
    db_badge = ('<div class="status-badge badge-ok">☁️ Cloud DB: Read/Write OK</div>' if st.session_state.get("db_verified")
                else '<div class="status-badge badge-warn">💾 DB: NOT saving</div>')
    c1, c2 = st.columns(2)
    c1.markdown(ai_badge, unsafe_allow_html=True)
    c2.markdown(db_badge, unsafe_allow_html=True)

    preset = st.selectbox("⚡ AI Model:", ["gemini-3.8-flash", "gemini-3.6-flash", "gemini-3.5-flash-lite"], index=0)
    custom_model = st.text_input("Custom model (optional):", placeholder="override the model name here")
    model_choice = custom_model.strip() or preset
    goal = st.number_input("🎯 Daily review goal:", min_value=1, max_value=50, value=DEFAULT_GOAL, key="daily_goal")
    st.divider()

    # ---- Quick capture: dump a thought, the AI files it under the right topic ----
    with st.expander("⚡ Quick Capture"):
        with st.form("qc_form", clear_on_submit=True):
            qc_text = st.text_area("Dump a thought, fact or idea:", height=90)
            qc_go = st.form_submit_button("Capture")
        if qc_go and qc_text.strip():
            qc_prompt = f"""You file notes into a personal knowledge base.
Existing topics: {", ".join(topic_names) or "(none)"}
Categories: {", ".join(CAT_OPTIONS)}
Return ONLY JSON: {{"topic": "exact existing title, or a short new title", "category": "one of the categories",
"tags": ["tag1", "tag2"], "note": "the thought rewritten as a clean, short note"}}
Thought: {qc_text}"""
            ok, res = ask_gemini(qc_prompt, selected_model=model_choice, json_mode=True)
            if ok:
                try:
                    d = parse_json_loose(res)
                    title = str(d.get("topic") or "Inbox").strip()
                    ex = next((t for t in st.session_state.topics if t["title"].lower() == title.lower()), None)
                    line = f"- ({today_str()}) {d.get('note') or qc_text}"
                    new_notes = ((ex.get("notes") if ex else "") + "\n" + line).strip()
                    tags = ", ".join(uniq(split_csv(ex.get("tags") if ex else "") + [str(x) for x in d.get("tags", [])]))
                    cat = d.get("category") if d.get("category") in CAT_OPTIONS else None
                    upsert_topic(title, category=None if ex else cat, notes=new_notes, tags=tags)
                    st.success(f"Filed under: {title}")
                    st.rerun()
                except Exception as e:
                    st.warning(f"Could not read the AI reply: {e}")
            else:
                st.error(res)

    st.markdown("#### 🔄 Add / Update Topic")
    mode = st.radio("Action:", ["➕ Create New", "✏️ Edit Existing"], horizontal=True, label_visibility="collapsed")

    if mode == "✏️ Edit Existing" and topic_names:
        sel_edit = st.selectbox("Select Topic:", topic_names)
        curr_t = next(t for t in st.session_state.topics if t["title"] == sel_edit)
        in_title = curr_t["title"]
        default_cat = curr_t.get("category") or CAT_OPTIONS[0]
        default_stat = curr_t.get("status") or "In Progress"
        default_mast = int(curr_t.get("mastery") or 50)
        existing_rels = [r for r in split_csv(curr_t.get("related_topics")) if r in topic_names]
        existing_pre = [r for r in split_csv(curr_t.get("prereq_topics")) if r in topic_names]
        default_tags = curr_t.get("tags") or ""
        wkey = f"edit_{sel_edit}_{default_mast}"
    else:
        in_title = st.text_input("Topic Title:", placeholder="e.g., 5G Site Rollout KPIs")
        default_cat, default_stat, default_mast = CAT_OPTIONS[0], "In Progress", 50
        existing_rels, existing_pre, default_tags = [], [], ""
        wkey = "new"

    in_cat = st.selectbox("Category:", CAT_OPTIONS,
                          index=CAT_OPTIONS.index(default_cat) if default_cat in CAT_OPTIONS else 0, key=f"cat_{wkey}")
    in_stat = st.selectbox("Status:", STAT_OPTIONS,
                           index=STAT_OPTIONS.index(default_stat) if default_stat in STAT_OPTIONS else 1, key=f"stat_{wkey}")
    in_mast = st.slider("Mastery (%):", 0, 100, default_mast, step=5, key=f"mast_{wkey}")

    others = [t for t in topic_names if t.lower() != in_title.strip().lower()]
    linked_topics = st.multiselect("🔗 Related topics:", others, default=existing_rels, key=f"links_{wkey}")
    prereq_sel = st.multiselect("⬆️ Prerequisites (learn these first):", others, default=existing_pre, key=f"pre_{wkey}")
    in_tags = st.text_input("🏷️ Tags (comma separated):", value=default_tags, key=f"tags_in_{wkey}_{default_tags}")

    if st.button("💾 Save / Update", type="primary"):
        if in_title.strip():
            act, err = upsert_topic(in_title, category=in_cat, status=in_stat, mastery=in_mast,
                                    related_topics=", ".join(linked_topics),
                                    prereq_topics=", ".join(prereq_sel), tags=in_tags)
            if not err:
                st.toast(f"☁️ Topic {act} & saved to cloud", icon="✅")
            st.rerun()
        else:
            st.warning("Enter a topic title.")

    if mode == "✏️ Edit Existing" and topic_names:
        if st.button("🗑️ Delete this topic"):
            delete_topic(in_title)
            st.rerun()

# ================= MAIN HEADER =================
st.markdown("""
<div class="main-header">
    <div class="eyebrow">PERSONAL KNOWLEDGE SYSTEM · AI-POWERED WORKSPACE</div>
    <h1>🧠 Nexus <span style="font-weight:500; opacity:.82;">/ Your Second Brain</span></h1>
    <p>Capture what you learn, connect ideas, review at the right time, and turn knowledge into clear presentations — all in one focused workspace.</p>
</div>
""", unsafe_allow_html=True)

if st.session_state.get("last_save_error"):
    st.error("⚠️ The last save did **NOT** reach the cloud:\n\n" + st.session_state.last_save_error)
elif not st.session_state.get("db_verified"):
    st.warning("⚠️ Data is currently **local only** and will be lost on refresh. Open the Cloud Diagnostics tab to see why.")

topic_list = [t["title"] for t in st.session_state.topics]
topics_all = st.session_state.topics

(tab_dash, tab_review, tab_map, tab_vault, tab_pitch,
 tab_coach, tab_ask, tab_search, tab_diag) = st.tabs([
    "📊 Dashboard", "🔁 Review", "🕸️ Visual Map", "🧠 Knowledge Vault", "🎯 Pitch Studio",
    "🎙️ Coach & Cards", "💬 Ask Your Brain", "🔎 Search", "⚙️ Cloud Diagnostics"
])

# ================= TAB: DASHBOARD =================
with tab_dash:
    if topics_all:
        df = pd.DataFrame(topics_all)
        due_now = due_topics()
        rt, streak = reviews_today(), compute_streak()

        m1, m2, m3, m4, m5, m6 = st.columns(6)
        m1.metric("📚 Topics", len(df))
        m2.metric("✅ Completed", int((df["status"] == "Completed").sum()))
        m3.metric("⚡ In Progress", int((df["status"] == "In Progress").sum()))
        m4.metric("🎯 Avg Mastery", f"{df['mastery'].mean():.0f}%")
        m5.metric("🔁 Due now", len(due_now))
        m6.metric("🔥 Streak", f"{streak} d")
        st.progress(min(rt / max(int(goal), 1), 1.0), text=f"Daily goal: {rt}/{int(goal)} reviews today")
        st.divider()

        ch1, ch2 = st.columns(2)
        with ch1:
            st.markdown("##### 📈 Mastery by Topic")
            st.bar_chart(df.set_index("title")["mastery"], color="#0284c7")
        with ch2:
            st.markdown("##### 📂 Topics per Domain")
            st.bar_chart(df["category"].value_counts(), color="#1e3a8a")

        hist = [a for a in st.session_state.activity if a.get("kind") == "mastery"]
        if hist:
            st.markdown("##### 🕒 Mastery over time")
            hdf = pd.DataFrame(hist)
            hdf["value"] = pd.to_numeric(hdf["value"], errors="coerce")
            pick = st.selectbox("Show:", ["All topics (average)"] + topic_list, key="hist_sel")
            if pick == "All topics (average)":
                piv = hdf.pivot_table(index="logged_on", columns="topic", values="value", aggfunc="last").sort_index().ffill()
                series = piv.mean(axis=1)
            else:
                series = hdf[hdf["topic"] == pick].groupby("logged_on")["value"].last()
            st.line_chart(series)

        weak = df[df["status"] != "Not Started"].sort_values("mastery").head(3)
        if not weak.empty:
            st.markdown("##### ⚠️ Weakest topics right now")
            st.write(" · ".join(f"**{r.title}** ({r.mastery}%)" for r in weak.itertuples()))

        st.markdown("### 📌 Knowledge Directory")
        f1, f2 = st.columns([2, 1])
        with f1:
            search_q = st.text_input("🔍 Filter by title or notes...", placeholder="Type keyword...", key="dir_search")
        with f2:
            cats = list(df["category"].unique())
            filter_cat = st.multiselect("Filter by Category:", cats, default=cats)
        fdf = df[df["category"].isin(filter_cat)]
        if search_q:
            fdf = fdf[fdf["title"].str.contains(search_q, case=False, na=False, regex=False) |
                      fdf["notes"].astype(str).str.contains(search_q, case=False, na=False, regex=False)]
        cols_to_show = ["title", "category", "status", "mastery", "tags", "next_review",
                        "prereq_topics", "related_topics", "updated_at", "notes"]
        st.dataframe(
            fdf[cols_to_show],
            column_config={
                "title": "Topic", "category": "Domain", "status": "Status",
                "mastery": st.column_config.ProgressColumn("Mastery", format="%d%%", min_value=0, max_value=100),
                "tags": "Tags", "next_review": "Next review", "prereq_topics": "⬆️ Prerequisites",
                "related_topics": "🔗 Related", "updated_at": "Updated", "notes": "Notes",
            },
            hide_index=True,
        )
        e1, e2 = st.columns(2)
        with e1:
            st.download_button("⬇️ Export filtered topics (CSV)",
                               data=fdf[cols_to_show].to_csv(index=False).encode("utf-8-sig"),
                               file_name="nexus_topics.csv", mime="text/csv")
        with e2:
            md_all = "\n\n---\n\n".join(
                f"# {t['title']}\n*{t.get('category') or ''} | {t.get('status') or ''} | {t.get('mastery') or 0}%*\n\n"
                f"Tags: {t.get('tags') or ''}\n\n{t.get('notes') or ''}\n\n{t.get('ai_summary') or ''}" for t in topics_all)
            st.download_button("⬇️ Export all notes (Markdown)", data=md_all.encode("utf-8"),
                               file_name="nexus_notes.md", mime="text/markdown")
    else:
        st.markdown('<div class="empty-state"><h3>Your knowledge base starts here</h3>'
                    '<p>Add a topic from the sidebar, or use Quick Capture to dump a thought and let the AI file it.</p></div>',
                    unsafe_allow_html=True)

# ================= TAB: REVIEW (SM-2 spaced repetition) =================
with tab_review:
    st.subheader("🔁 Spaced Repetition Review")
    due_list = due_topics()
    if not due_list:
        st.success("Nothing is due right now 🎉 Come back later, or add more topics.")
    else:
        t = due_list[0]
        st.caption(f"{len(due_list)} topic(s) due · weakest first")
        st.markdown(f"## {t['title']}")
        st.markdown(f"`{t.get('category')}` · mastery **{t.get('mastery')}%** · "
                    f"streak of good reviews: **{t.get('repetitions')}** · interval **{t.get('interval_days')} d**")
        st.info("Try to recall the key points from memory first, then reveal your notes and rate yourself honestly.")

        if st.session_state.get("rev_show_for") != t["title"]:
            if st.button("👀 Reveal notes", key="rev_reveal", type="primary"):
                st.session_state.rev_show_for = t["title"]
                st.rerun()
        else:
            with st.container(border=True):
                st.markdown(render_links(t.get("notes") or "_No notes yet. Add some in the Knowledge Vault._"))
                if t.get("ai_summary"):
                    with st.expander("🤖 AI summary"):
                        st.markdown(t["ai_summary"])
                deck = [c for c in st.session_state.flashcards if c["topic"] == t["title"]]
                if deck:
                    with st.expander(f"🃏 {len(deck)} flashcards for this topic"):
                        for c in deck:
                            st.markdown(f"**Q:** {c['q']}\n\n**A:** {c['a']}")
            st.markdown("**How well did you remember it?**")
            rcols = st.columns(4)
            for col, (label, g) in zip(rcols, GRADES.items()):
                if col.button(label, key=f"grade_{g}"):
                    nxt = sm2_next(t, g)
                    new_m = max(0, min(100, int(t.get("mastery") or 0) + MASTERY_DELTA[g]))
                    upsert_topic(t["title"], mastery=new_m, **nxt)
                    log_activity(t["title"], "review", g)
                    st.session_state.rev_show_for = None
                    st.rerun()

    upcoming = sorted([x for x in topics_all if x.get("next_review") and x["next_review"] > today_str()],
                      key=lambda x: x["next_review"])[:10]
    if upcoming:
        st.markdown("##### 📅 Upcoming reviews")
        st.dataframe(pd.DataFrame(upcoming)[["title", "next_review", "mastery", "interval_days"]],
                     column_config={"title": "Topic", "next_review": "Next review", "mastery": "Mastery %",
                                    "interval_days": "Interval (days)"}, hide_index=True)

# ================= TAB: VISUAL MAP =================
with tab_map:
    st.subheader("🕸️ Interactive Second-Brain Knowledge Graph")
    if not topics_all:
        st.markdown('<div class="empty-state"><h3>No nodes to visualize yet</h3>'
                    '<p>Create a few topics first, then use AI Auto-Connect.</p></div>', unsafe_allow_html=True)
    else:
        fc1, fc2, fc3 = st.columns([2, 2, 1])
        all_cats = sorted({t.get("category") or "General" for t in topics_all})
        sel_cats = fc1.multiselect("Domains:", all_cats, default=all_cats)
        sel_stats = fc2.multiselect("Status:", STAT_OPTIONS, default=STAT_OPTIONS)
        engine = fc3.radio("Graph engine:", ["Interactive", "Static"], horizontal=True)
        vis = [t for t in topics_all if (t.get("category") or "General") in sel_cats and t.get("status") in sel_stats]

        g1, g2 = st.columns([3, 1])
        with g2:
            show_categories = st.checkbox("Show domain hubs", value=True)
            physics_enabled = st.checkbox("Bouncy physics", value=True)
            highlight_weak = st.checkbox("Highlight weak topics (<40%)", value=False)

            if st.button("🤖 AI Auto-Connect", type="primary", key="auto_connect"):
                if len(topics_all) >= 2:
                    with st.spinner("Discovering topic connections..."):
                        t_list_str = "\n".join(f"- {t['title']} ({t.get('category', '')})" for t in topics_all)
                        link_prompt = f"""Analyze the relationships between these topics:
{t_list_str}
Return ONLY valid JSON where every key is an exact topic title:
{{"Topic A": {{"related": ["Topic B"], "prereqs": ["Topic C"]}}}}
"prereqs" are topics that should be learned BEFORE Topic A. Use exact titles from the list only."""
                        ok, res = ask_gemini(link_prompt, selected_model=model_choice, json_mode=True)
                        if ok:
                            linked_ok = False
                            try:
                                links = normalize_links(parse_json_loose(res))
                                known = {t["title"].lower(): t["title"] for t in topics_all}
                                for t_title, rel in links.items():
                                    me = known.get(t_title.lower())
                                    if not me:
                                        continue
                                    cur = next(t for t in st.session_state.topics if t["title"] == me)
                                    new_rel = [known[x.lower()] for x in rel["related"] if x.lower() in known and known[x.lower()] != me]
                                    new_pre = [known[x.lower()] for x in rel["prereqs"] if x.lower() in known and known[x.lower()] != me]
                                    upsert_topic(me,
                                                 related_topics=", ".join(uniq(split_csv(cur.get("related_topics")) + new_rel)),
                                                 prereq_topics=", ".join(uniq(split_csv(cur.get("prereq_topics")) + new_pre)))
                                linked_ok = True
                            except Exception as e:
                                st.warning(f"Could not read the AI reply: {e}")
                                st.code(res[:500])
                            if linked_ok:
                                st.rerun()
                        else:
                            st.error(res)
                else:
                    st.info("Add at least 2 topics first.")

            if st.button("🧭 Build learning path", key="lp_btn"):
                st.session_state.show_path = True
            st.markdown("**🎨 Legend:**\n- 🟦 Domain hub\n- 🟢 Strong (≥80%)\n- 🟠 In progress\n- 🔴 Early stage\n"
                        "- 🟧 Orange arrow: prerequisite → topic\n- 🔵 Blue line: related")

        with g1:
            if not vis:
                st.info("No topics match the current filters.")
            elif engine == "Interactive" and AGRAPH_INSTALLED:
                nodes, edges, added, seen_edges = [], [], set(), set()
                vis_titles = {t["title"] for t in vis}
                if show_categories:
                    for c in set(t.get("category") or "General" for t in vis):
                        nodes.append(Node(id=f"HUB_{c}", label=f"📂 {c}", size=26, color="#1e3a8a",
                                          shape="hexagon", font={"color": "#ffffff"}))
                        added.add(f"HUB_{c}")
                for t in vis:
                    mast = int(t.get("mastery") or 50)
                    if t["title"] not in added:
                        nodes.append(Node(id=t["title"], label=f"{t['title']}\n({mast}%)", size=14 + int(mast / 6),
                                          color=node_color(t, highlight_weak), shape="dot", font={"color": "#e8eefb"}))
                        added.add(t["title"])
                    if show_categories:
                        edges.append(Edge(source=f"HUB_{t.get('category') or 'General'}", target=t["title"], color="#94a3b8"))
                for t in vis:
                    for r in split_csv(t.get("related_topics")):
                        tgt = next((x for x in vis_titles if x.lower() == r.lower()), None)
                        key = frozenset((t["title"], tgt or ""))
                        if tgt and tgt != t["title"] and key not in seen_edges:
                            seen_edges.add(key)
                            edges.append(Edge(source=t["title"], target=tgt, color="#38bdf8", width=2))
                    for p in split_csv(t.get("prereq_topics")):
                        src = next((x for x in vis_titles if x.lower() == p.lower()), None)
                        if src and src != t["title"]:
                            edges.append(Edge(source=src, target=t["title"], color="#f59e0b", width=3, arrows="to"))
                try:
                    clicked_node = agraph(nodes=nodes, edges=edges,
                                          config=Config(width=900, height=540, directed=False, physics=physics_enabled))
                    if clicked_node and clicked_node != st.session_state.get("agraph_prev"):
                        st.session_state.agraph_prev = clicked_node
                        st.session_state.graph_clicked = clicked_node
                except Exception as e:
                    st.warning(f"Interactive graph failed ({str(e)[:120]}). Showing a static graph instead.")
                    st.graphviz_chart(build_dot(vis, show_categories, highlight_weak))
            else:
                if engine == "Interactive":
                    st.info("streamlit-agraph is not installed, showing the static graph.")
                st.graphviz_chart(build_dot(vis, show_categories, highlight_weak))

        if st.session_state.get("show_path"):
            st.markdown("#### 🧭 Suggested learning path")
            st.caption("Prerequisites first; among topics that are ready, the weakest comes first.")
            path_rows = learning_path(topics_all)
            st.dataframe(pd.DataFrame(path_rows),
                         column_config={"mastery": st.column_config.ProgressColumn("Mastery", format="%d%%", min_value=0, max_value=100)},
                         hide_index=True)
            if st.button("🧠 Turn this path into a study plan (AI)", key="lp_ai"):
                with st.spinner("Planning..."):
                    plan_prompt = ("Write a short, practical study plan (max 12 lines) for this ordered list of topics. "
                                   "Spend more time on low mastery topics.\n" +
                                   "\n".join(f"{r['#']}. {r['topic']} (mastery {r['mastery']}%, {r['note']})" for r in path_rows))
                    ok, plan = ask_gemini(plan_prompt, selected_model=model_choice)
                    st.markdown(plan) if ok else st.error(plan)

        # ---- Topic details panel (click a node, or pick here) ----
        st.divider()
        st.markdown("### 📝 Topic Details")
        names_g = [t["title"] for t in topics_all]
        clicked = st.session_state.get("graph_clicked")
        if clicked and str(clicked).startswith("HUB_"):
            cat_name = str(clicked)[4:]
            in_cat_list = [t for t in topics_all if (t.get("category") or "General") == cat_name]
            st.markdown(f"**📂 {cat_name}** ({len(in_cat_list)} topics)")
            if in_cat_list:
                st.dataframe(pd.DataFrame(in_cat_list)[["title", "status", "mastery"]], hide_index=True)
            clicked = None
        start_idx = names_g.index(clicked) if clicked in names_g else 0
        sel_g = st.selectbox("Open topic (click a node in the graph or pick here):", names_g, index=start_idx)
        tg = next(t for t in topics_all if t["title"] == sel_g)

        d1, d2, d3, d4 = st.columns(4)
        d1.metric("Status", tg.get("status") or "-")
        d2.metric("Mastery", f"{int(tg.get('mastery') or 0)}%")
        d3.metric("Next review", tg.get("next_review") or "-")
        d4.metric("Updated", tg.get("updated_at") or "-")
        if tg.get("tags"):
            st.markdown("🏷️ " + " ".join(f"`{x}`" for x in split_csv(tg["tags"])))
        if tg.get("prereq_topics"):
            st.markdown(f"**⬆️ Learn first:** {tg['prereq_topics']}")
        if tg.get("related_topics"):
            st.markdown(f"**🔗 Related:** {tg['related_topics']}")

        nav = uniq([x for x in extract_links(tg.get("notes")) if x.lower() in [n.lower() for n in names_g]]
                   + backlinks_of(sel_g))
        if nav:
            st.markdown("**↪️ Jump to a linked note:**")
            ncols = st.columns(min(len(nav), 4))
            for i, name in enumerate(nav):
                real = next(n for n in names_g if n.lower() == name.lower())
                if ncols[i % len(ncols)].button(real, key=f"nav_{sel_g}_{real}"):
                    st.session_state.graph_clicked = real
                    st.rerun()
            back = backlinks_of(sel_g)
            if back:
                st.caption("Mentioned in: " + ", ".join(back))

        g_notes = st.text_area("Notes", value=tg.get("notes") or "", height=200,
                               key=f"graph_notes_{sel_g}_{abs(hash(tg.get('notes') or ''))}")
        st.caption("Tip: write [[Another Topic]] inside notes to link topics together.")
        if st.button("💾 Save notes", key=f"graph_save_{sel_g}"):
            _, err = upsert_topic(sel_g, notes=g_notes)
            if err:
                st.error(err)
            else:
                st.success("Notes saved to cloud ✅")
        if tg.get("ai_summary"):
            with st.expander("🤖 AI Summary"):
                st.markdown(tg["ai_summary"])

# ================= TAB: KNOWLEDGE VAULT =================
with tab_vault:
    st.subheader("📥 Multi-Modal Knowledge Ingestion")
    if topic_list:
        active_name = st.selectbox("Select Topic:", topic_list)
        active = next(t for t in st.session_state.topics if t["title"] == active_name)

        v1, v2 = st.columns(2)
        with v1:
            user_notes = st.text_area("✍️ Notes & Key Formulas (use [[Topic]] to link):", value=active.get("notes") or "",
                                      height=200, key=f"notes_{active_name}_{abs(hash(active.get('notes') or ''))}")
            tags_val = active.get("tags") or ""
            user_tags = st.text_input("🏷️ Tags:", value=tags_val, key=f"vtags_{active_name}_{abs(hash(tags_val))}")
            b1, b2 = st.columns(2)
            if b1.button("💾 Save notes & tags"):
                _, err = upsert_topic(active_name, notes=user_notes, tags=user_tags)
                st.error(err) if err else st.success("Saved to cloud ✅")
            if b2.button("🏷️ Suggest tags (AI)"):
                ok, res = ask_gemini(f"Suggest 3 to 6 short lowercase tags for this topic. Return ONLY a JSON array of strings.\n"
                                     f"Title: {active_name}\nNotes: {user_notes[:2000]}", selected_model=model_choice, json_mode=True)
                if ok:
                    try:
                        new_tags = [str(x).strip().lower() for x in parse_json_loose(res)]
                        upsert_topic(active_name, tags=", ".join(uniq(split_csv(user_tags) + new_tags)))
                        st.rerun()
                    except Exception as e:
                        st.warning(f"Could not read the AI reply: {e}")
                else:
                    st.error(res)
            with st.expander("👁️ Preview with links"):
                st.markdown(render_links(user_notes) or "_Empty_")
                back = backlinks_of(active_name)
                if back:
                    st.caption("Linked from: " + ", ".join(back))

        with v2:
            up_files = st.file_uploader("📄 PDFs or slide images (several allowed):",
                                        type=["pdf", "png", "jpg", "jpeg"], accept_multiple_files=True)
            url_in = st.text_input("🌐 Web page or YouTube link (optional):", placeholder="https://...")
            extract_glossary = st.checkbox("Auto-extract terms to Glossary", value=True)

            if st.button("⚡ AI Summarize", type="primary"):
                if not up_files and not user_notes and not url_in.strip():
                    st.warning("Add a file, a link, or some notes first.")
                else:
                    with st.spinner(f"Analyzing with {model_choice}..."):
                        text_parts, media = [], []
                        for f in up_files or []:
                            if f.name.lower().endswith(".pdf"):
                                txt, part = pdf_to_input(f)
                                if txt:
                                    text_parts.append(f"[{f.name}]\n{txt}")
                                if part is not None:
                                    media.append(part)
                                    st.caption(f"{f.name} looks scanned, sending it to the AI to read directly.")
                            else:
                                try:
                                    media.append(Image.open(f))
                                except Exception as e:
                                    st.error(f"Could not read {f.name}: {e}")
                        if url_in.strip():
                            ok_u, url_text = fetch_url_text(url_in)
                            if ok_u:
                                text_parts.append(f"[{url_in}]\n{url_text}")
                            else:
                                st.warning(url_text)
                        content = "\n\n".join(text_parts)[:14000]
                        prompt = f"""Summarize this material for the topic '{active_name}'.
Notes: {user_notes}
Content: {content}

Format:
### 📌 Executive Summary (4 bullets)
### ⚙️ Key Takeaways
### 📖 Key Terms (Strictly 3 lines: TERM: [name] | DEF: [short definition])
Last line only: TAGS: tag1, tag2, tag3"""
                        ok, ai_res = ask_gemini(prompt, media=media, selected_model=model_choice)
                        if ok:
                            new_tags = []
                            for line in ai_res.splitlines():
                                if line.strip().upper().startswith("TAGS:"):
                                    new_tags = [x.strip().lower() for x in line.split(":", 1)[1].split(",") if x.strip()]
                            body = "\n".join(l for l in ai_res.splitlines() if not l.strip().upper().startswith("TAGS:"))
                            upsert_topic(active_name, notes=user_notes, ai_summary=body,
                                         tags=", ".join(uniq(split_csv(user_tags) + new_tags)))
                            if extract_glossary:
                                for line in body.split("\n"):
                                    if "TERM:" in line and "DEF:" in line:
                                        term_part, def_part = line.split("DEF:", 1)
                                        term = term_part.replace("TERM:", "").replace("*", "").replace("|", "").strip(" -•:[]")
                                        if term:
                                            upsert_glossary(term, "Technical", def_part.strip(" *[]"))
                            st.rerun()
                        else:
                            st.error(ai_res)

        if active.get("ai_summary"):
            st.divider()
            st.markdown("#### 🤖 Saved AI Synthesis")
            with st.container(border=True):
                st.markdown(active["ai_summary"])
    else:
        st.info("Add a topic from the sidebar first.")

# ================= TAB 4: PITCH STUDIO =================
with tab_pitch:
    st.subheader("🎯 Automated Executive Presentation Studio")
    p1, p2, p3 = st.columns(3)
    with p1:
        pitch_topic = st.selectbox("Select Topic:", topic_list if topic_list else ["General"], key="pitch_sel")
    with p2:
        audience = st.selectbox("Target Audience:", ["Executive Board / PMO", "Engineering Deep-Dive", "Oral Assessment Panel"])
    with p3:
        slide_count = st.slider("Slides:", 3, 6, 4)

    if st.button("✨ Generate Slide Deck + Script", type="primary"):
        t_obj = next((t for t in st.session_state.topics if t['title'] == pitch_topic), {})
        with st.spinner("Generating presentation..."):
            pitch_prompt = f"""Create a {slide_count}-slide presentation on '{pitch_topic}' for '{audience}'.
Notes: {(t_obj.get('notes') or '')[:1500]}
Summary: {(t_obj.get('ai_summary') or '')[:1500]}
Return ONLY a JSON array:
[{{"title": "Slide Title", "bullets": ["Point 1", "Point 2"], "speaker_notes": "Script...", "time": "2 mins"}}]"""
            ok, raw = ask_gemini(pitch_prompt, selected_model=model_choice, json_mode=True)
            if ok:
                try:
                    slides_data = parse_json_loose(raw)
                    st.session_state.last_deck = (pitch_topic, audience, slides_data)
                except Exception:
                    st.write(raw)
            else:
                st.error(raw)

    if st.session_state.get("last_deck"):
        d_topic, d_aud, slides_data = st.session_state.last_deck
        cols = st.columns(2)
        for i, s in enumerate(slides_data):
            with cols[i % 2]:
                with st.container(border=True):
                    st.markdown(f"#### 📊 Slide {i+1}: {s.get('title')} `{s.get('time')}`")
                    for b in s.get("bullets", []):
                        st.markdown(f"- {b}")
                    with st.expander("🎙️ Speaker Script"):
                        st.write(s.get("speaker_notes"))
        st.download_button("📥 Download PowerPoint (.pptx)",
                           data=build_advanced_pptx(d_topic, d_aud, slides_data),
                           file_name=f"{d_topic}_Deck.pptx",
                           mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
                           type="primary")

# ================= TAB: COACH & CARDS =================
with tab_coach:
    c_left, c_right = st.columns([3, 2])
    with c_left:
        st.subheader("🎙️ Active Recall Coach")
        coach_mode = st.radio("Mode:", ["🧠 Feynman", "🔍 Gap Detector", "🃏 Flashcards", "📝 Quiz"], horizontal=True)
        coach_topic = st.selectbox("Topic:", topic_list if topic_list else ["General"], key="coach_sel")
        c_obj = next((t for t in st.session_state.topics if t["title"] == coach_topic), {})

        if coach_mode == "🧠 Feynman":
            st.caption("Explain the topic in your own words, typed or recorded. The AI scores it and polishes it.")
            user_exp = st.text_area("Your explanation:", height=120)
            audio = st.audio_input("Or record it (optional)") if hasattr(st, "audio_input") else None
            if st.button("🔍 Evaluate", type="primary") and (user_exp or audio):
                with st.spinner("Evaluating..."):
                    media = [types.Part.from_bytes(data=audio.getvalue(), mime_type="audio/wav")] if audio else []
                    prompt = (f"The student explains '{coach_topic}'.\n"
                              + (f"Typed explanation: {user_exp}\n" if user_exp else "")
                              + ("The explanation is also in the attached audio: transcribe it first, then evaluate it.\n" if audio else "")
                              + "Reference notes: " + (c_obj.get("notes") or "")[:1500] + "\n"
                              "Give: transcript (if audio), score /10, missing key terms, mistakes, and a polished 3-sentence version.")
                    ok, res = ask_gemini(prompt, media=media, selected_model=model_choice)
                    st.markdown(res) if ok else st.error(res)

        elif coach_mode == "🔍 Gap Detector":
            st.caption("Compares your notes against what a student should know, and lists what is missing.")
            if st.button("🔎 Find my knowledge gaps", type="primary"):
                with st.spinner("Analyzing..."):
                    prompt = (f"Topic: {coach_topic}\nMy notes:\n{(c_obj.get('notes') or '')[:3000]}\n"
                              f"My summary:\n{(c_obj.get('ai_summary') or '')[:2000]}\n\n"
                              "List the key sub-topics someone should know about this topic. For each, mark COVERED, PARTIAL or MISSING "
                              "based on my notes. Then give 5 questions that would fill the biggest gaps. Use short markdown.")
                    ok, res = ask_gemini(prompt, selected_model=model_choice)
                    st.markdown(res) if ok else st.error(res)

        elif coach_mode == "🃏 Flashcards":
            if st.button("⚡ Generate 6 flashcards", type="primary"):
                with st.spinner("Creating cards..."):
                    prompt = (f"Return ONLY a JSON array of 6 flashcards for '{coach_topic}' based on these notes: "
                              f"{(c_obj.get('notes') or '')[:2500]} {(c_obj.get('ai_summary') or '')[:1500]} "
                              'Format: [{"q":"Question?","a":"Answer"}]')
                    ok, res = ask_gemini(prompt, selected_model=model_choice, json_mode=True)
                    if ok:
                        try:
                            st.session_state.pending_cards = {"topic": coach_topic, "cards": parse_json_loose(res)}
                        except Exception:
                            st.write(res)
                    else:
                        st.error(res)
            pend = st.session_state.get("pending_cards")
            if pend:
                st.markdown(f"**New cards for {pend['topic']}** (not saved yet)")
                for i, fc in enumerate(pend["cards"]):
                    st.markdown(f'<div class="flashcard"><b>Q{i+1}: {fc.get("q")}</b><br><span style="opacity:.8">{fc.get("a")}</span></div>',
                                unsafe_allow_html=True)
                s1, s2 = st.columns(2)
                if s1.button("💾 Save all to my deck"):
                    err = add_flashcards(pend["topic"], pend["cards"])
                    st.session_state.pending_cards = None
                    if err:
                        st.error(err)
                    else:
                        st.toast("Cards saved ✅")
                    st.rerun()
                if s2.button("🗑️ Discard"):
                    st.session_state.pending_cards = None
                    st.rerun()

            st.divider()
            st.markdown("#### 📚 My deck")
            f_topics = ["All topics"] + sorted({c["topic"] for c in st.session_state.flashcards})
            f_pick = st.selectbox("Deck:", f_topics, key="deck_pick")
            deck = [c for c in st.session_state.flashcards if f_pick == "All topics" or c["topic"] == f_pick]
            if not deck:
                st.caption("No saved cards yet. Generate some and save them.")
            else:
                st.caption(f"{len(deck)} cards")
                study = st.toggle("Study mode", value=False)
                if study:
                    idx = st.session_state.get("study_idx", 0) % len(deck)
                    card = deck[idx]
                    st.markdown(f'<div class="flashcard"><b>{card["topic"]}</b><br>Q: {card["q"]}</div>', unsafe_allow_html=True)
                    if not st.session_state.get("study_show"):
                        if st.button("👀 Show answer", key="study_reveal"):
                            st.session_state.study_show = True
                            st.rerun()
                    else:
                        st.success(card["a"])
                        k1, k2 = st.columns(2)
                        for col, label, val in ((k1, "✅ Got it", 1), (k2, "❌ Again", 0)):
                            if col.button(label, key=f"study_{val}"):
                                log_activity(card["topic"], "flashcard", val)
                                st.session_state.study_idx = idx + 1
                                st.session_state.study_show = False
                                st.rerun()
                else:
                    for card in deck:
                        with st.expander(f"{card['topic']}: {card['q'][:70]}"):
                            st.write(card["a"])
                            if st.button("Delete card", key=f"delcard_{card['topic']}_{abs(hash(card['q']))}"):
                                delete_flashcard(card["topic"], card["q"])
                                st.rerun()
                anki = pd.DataFrame([{"Front": c["q"], "Back": c["a"], "Tags": c["topic"].replace(" ", "_")} for c in deck])
                st.download_button("⬇️ Export deck for Anki (CSV)", data=anki.to_csv(index=False).encode("utf-8-sig"),
                                   file_name="nexus_flashcards.csv", mime="text/csv")

        else:
            if st.button("🎲 Generate 5 questions", type="primary"):
                with st.spinner("Generating quiz..."):
                    quiz_prompt = f"""Create 5 multiple-choice questions on '{coach_topic}'.
Base them on these notes and summary when present.
Notes: {(c_obj.get('notes') or '')[:2000]}
Summary: {(c_obj.get('ai_summary') or '')[:2000]}
Return ONLY a JSON array like:
[{{"q": "Question?", "options": ["A", "B", "C", "D"], "answer": 0, "why": "short explanation"}}]
where "answer" is the 0-based index of the correct option."""
                    ok, res = ask_gemini(quiz_prompt, selected_model=model_choice, json_mode=True)
                    if ok:
                        try:
                            items = parse_json_loose(res)
                            for i in range(20):
                                st.session_state.pop(f"quiz_{i}", None)
                            st.session_state.quiz = {"topic": coach_topic, "items": items, "result": None}
                        except Exception:
                            st.write(res)
                    else:
                        st.error(res)

            qz = st.session_state.get("quiz")
            if qz and qz["topic"] == coach_topic:
                for i, it in enumerate(qz["items"]):
                    st.markdown(f"**Q{i+1}. {it.get('q')}**")
                    st.radio("Answer:", it.get("options", []), index=None, key=f"quiz_{i}", label_visibility="collapsed")
                if st.button("✅ Submit answers"):
                    score, review = 0, []
                    for i, it in enumerate(qz["items"]):
                        opts = it.get("options", [])
                        try:
                            correct = opts[int(it.get("answer", 0))]
                        except Exception:
                            correct = None
                        chosen = st.session_state.get(f"quiz_{i}")
                        is_right = chosen is not None and chosen == correct
                        score += int(is_right)
                        review.append((i + 1, is_right, correct, it.get("why", "")))
                    pct = int(score / max(len(qz["items"]), 1) * 100)
                    old_m, new_m = None, None
                    if c_obj:
                        old_m = int(c_obj.get("mastery") or 50)
                        new_m = max(0, min(100, int(round((old_m + pct) / 2 / 5) * 5)))
                        upsert_topic(coach_topic, mastery=new_m)
                    log_activity(coach_topic, "quiz", pct)
                    qz["result"] = {"score": score, "total": len(qz["items"]), "review": review, "old": old_m, "new": new_m}
                res_q = qz.get("result")
                if res_q:
                    st.metric("Score", f"{res_q['score']}/{res_q['total']}")
                    if res_q["new"] is not None:
                        st.info(f"Mastery updated: {res_q['old']}% -> {res_q['new']}%")
                    for num, ok_q, correct, why in res_q["review"]:
                        if ok_q:
                            st.success(f"Q{num}: correct. {why}")
                        else:
                            st.error(f"Q{num}: the correct answer is: {correct}. {why}")

    with c_right:
        st.subheader("📖 Smart Glossary")
        with st.expander("➕ Add Term"):
            g_t = st.text_input("Term:")
            g_c = st.selectbox("Domain:", ["PM", "Technical", "Business"])
            g_d = st.text_area("Definition:")
            if st.button("Save Term") and g_t.strip():
                err = upsert_glossary(g_t, g_c, g_d)
                if not err:
                    st.toast("Term saved ✅")
                st.rerun()
        if st.session_state.glossary:
            gl_q = st.text_input("🔍 Search glossary:", key="gl_search")
            gdf = pd.DataFrame(st.session_state.glossary)
            show = [c for c in ["term", "category", "context"] if c in gdf.columns]
            if gl_q:
                mask = (gdf["term"].astype(str).str.contains(gl_q, case=False, na=False, regex=False) |
                        gdf["context"].astype(str).str.contains(gl_q, case=False, na=False, regex=False))
                gdf = gdf[mask]
            st.dataframe(gdf[show], hide_index=True)
            with st.expander("🗑️ Delete a term"):
                del_term = st.selectbox("Term:", [g["term"] for g in st.session_state.glossary], key="gl_del")
                if st.button("Delete term"):
                    delete_glossary(del_term)
                    st.rerun()

# ================= TAB: ASK YOUR BRAIN =================
with tab_ask:
    st.subheader("💬 Ask Your Brain")
    st.caption("Answers are based only on your own topics, notes and summaries.")
    if "chat" not in st.session_state:
        st.session_state.chat = []
    if not topics_all:
        st.info("Add some topics and notes first.")
    else:
        if st.session_state.chat and st.button("🧹 Clear chat"):
            st.session_state.chat = []
            st.rerun()
        for role, msg in st.session_state.chat:
            with st.chat_message(role):
                st.markdown(msg)
        user_q = st.chat_input("Ask something about your knowledge base...")
        if user_q:
            ctx = "\n\n".join(
                f"## {t['title']} ({t.get('category') or ''}, {t.get('status') or ''}, mastery {t.get('mastery') or 0}%, tags: {t.get('tags') or ''})\n"
                f"Notes: {(t.get('notes') or '')[:800]}\nSummary: {(t.get('ai_summary') or '')[:800]}"
                for t in topics_all)[:20000]
            hist_txt = "\n".join(f"{r}: {m}" for r, m in st.session_state.chat[-6:])
            ask_prompt = f"""You are the user's second brain. Answer ONLY from the knowledge base below and mention which topics you used.
If the answer is not in it, say so and suggest what to add.

KNOWLEDGE BASE:
{ctx}

CHAT SO FAR:
{hist_txt}

QUESTION: {user_q}"""
            with st.spinner("Thinking..."):
                ok, ans = ask_gemini(ask_prompt, selected_model=model_choice)
            st.session_state.chat.append(("user", user_q))
            st.session_state.chat.append(("assistant", ans))
            st.rerun()

# ================= TAB: SEARCH =================
with tab_search:
    st.subheader("🔎 Search everything")
    all_tags = sorted({x for t in topics_all for x in split_csv(t.get("tags"))})
    s1, s2 = st.columns([3, 2])
    q_all = s1.text_input("Search topics, notes, summaries, glossary and flashcards:", key="global_q")
    tag_sel = s2.multiselect("Filter by tags:", all_tags)
    if q_all or tag_sel:
        hits = []
        for t in topics_all:
            tags_t = [x.lower() for x in split_csv(t.get("tags"))]
            if tag_sel and not all(x.lower() in tags_t for x in tag_sel):
                continue
            hay = " ".join([t["title"], t.get("notes") or "", t.get("ai_summary") or "", t.get("tags") or ""])
            if not q_all or q_all.lower() in hay.lower():
                hits.append(t)
        st.markdown(f"**{len(hits)} topic(s)**")
        for t in hits[:30]:
            with st.container(border=True):
                st.markdown(f"**{t['title']}** · {t.get('status')} · {t.get('mastery')}%  " +
                            " ".join(f"`{x}`" for x in split_csv(t.get("tags"))))
                if q_all:
                    src = (t.get("notes") or "") if q_all.lower() in (t.get("notes") or "").lower() else (t.get("ai_summary") or "")
                    st.caption(snippet(src, q_all))
        if q_all and not tag_sel:
            gl_hits = [g for g in st.session_state.glossary
                       if q_all.lower() in f"{g.get('term', '')} {g.get('context', '')}".lower()]
            fc_hits = [c for c in st.session_state.flashcards
                       if q_all.lower() in f"{c['q']} {c['a']}".lower()]
            if gl_hits:
                with st.expander(f"📖 {len(gl_hits)} glossary match(es)"):
                    for g in gl_hits[:20]:
                        st.markdown(f"**{g['term']}**: {g.get('context', '')}")
            if fc_hits:
                with st.expander(f"🃏 {len(fc_hits)} flashcard match(es)"):
                    for c in fc_hits[:20]:
                        st.markdown(f"**{c['q']}**\n\n{c['a']}")
    elif all_tags:
        st.markdown("**Tags:** " + " ".join(f"`{x}`" for x in all_tags))
    else:
        st.caption("Type something to search, or add tags to your topics.")

# ================= TAB: CLOUD DIAGNOSTICS =================
with tab_diag:
    st.subheader("⚙️ Cloud Database Diagnostics & Backup")
    st.write(f"**Supabase API URL:** `{SUPABASE_URL or 'Not Set'}`")
    st.write(f"**Supabase key set:** `{'yes' if SUPABASE_KEY else 'no'}`  |  **Mode:** `{'per-user login' if USE_AUTH else 'shared (no login)'}`  |  "
             f"**Topics:** `{len(st.session_state.topics)}`  |  **Cards:** `{len(st.session_state.flashcards)}`")

    if st.session_state.get("db_verified"):
        st.success("✅ Supabase read and write both work.")
    else:
        st.error("❌ Cloud saving is not working. Reason:")
        st.code(st.session_state.get("db_error_msg") or "Unknown")

    if st.button("🔄 Run full read/write test & reload from cloud", type="primary"):
        ok_, msg_ = run_db_check()
        st.session_state.db_verified = ok_
        st.session_state.db_error_msg = msg_
        if ok_:
            load_data_from_cloud()
            st.session_state.last_save_error = ""
        st.rerun()

    with st.expander("🧱 SQL setup: run this once in Supabase > SQL Editor", expanded=not st.session_state.get("db_verified")):
        st.code(SQL_SETUP, language="sql")
    with st.expander("🔐 Optional: per-user login (set USE_AUTH = \"true\" in Secrets, then run this AFTER the setup SQL)"):
        st.code(SQL_AUTH, language="sql")

    st.divider()
    st.markdown("#### 📦 Manual JSON Backup & Restore")
    backup_data = json.dumps({"topics": st.session_state.topics, "glossary": st.session_state.glossary,
                              "flashcards": st.session_state.flashcards, "activity": st.session_state.activity},
                             indent=2, default=str)
    st.download_button("⬇️ Download Full Backup (.json)", data=backup_data,
                       file_name="nexus_backup.json", mime="application/json")

    restore = st.file_uploader("Restore from backup (.json)", type=["json"], key="restore_up")
    if restore and st.button("♻️ Restore into app (and cloud)"):
        try:
            data = json.load(restore)
            for t in data.get("topics", []):
                fields = {k: v for k, v in t.items() if k in TOPIC_DEFAULTS and k not in ("title", "updated_at")}
                upsert_topic(t.get("title", "Untitled"), **fields)
            for g in data.get("glossary", []):
                upsert_glossary(g.get("term", ""), g.get("category"), g.get("context"))
            by_topic = {}
            for c in data.get("flashcards", []):
                by_topic.setdefault(c.get("topic", ""), []).append({"q": c.get("q"), "a": c.get("a")})
            for tp, cards in by_topic.items():
                add_flashcards(tp, cards)
            old_act = data.get("activity", [])
            if old_act:
                clean_act = [{"topic": a.get("topic"), "kind": a.get("kind"), "value": a.get("value"),
                              "logged_on": a.get("logged_on")} for a in old_act]
                st.session_state.activity.extend(clean_act)
                _cloud_insert("activity_log", clean_act)
            st.success("Restored ✅")
            st.rerun()
        except Exception as e:
            st.error(f"Restore failed: {e}")
